"""
tracit_check.py - collects what's needed to automate the TracIT EUC report
download, WITHOUT copying passwords, cookies, tokens, names or serials.

Run it the same way you run the app:
    python tracit_check.py

It looks on your Desktop (and in Downloads) for:
  - the newest .har file  (Edge F12 -> Network -> Export HAR)
  - the newest EUC report (.xlsx / .csv with "euc" in the name)
and writes TracIT_Export_Details.txt to your Desktop. Nothing is sent
anywhere except the one download test in step 2, which repeats the same
export request to TracIT using your Windows login (like the app would).
Add --skip-download-test to leave that out.
"""
import glob
import json
import os
import re
import sys
import tempfile
from datetime import datetime

SECRET = re.compile(r"cookie|authorization|token|xsrf|csrf|session|secret|passw|apikey|api-key|"
                    r"signature|^sig$|^code$|bearer|saml|jwt|auth", re.I)
STATIC = re.compile(r"\.(js|css|png|jpe?g|gif|svg|woff2?|ttf|ico|map)(\?|$)", re.I)
lines = []


def out(text=""):
    lines.append(text)
    print(text)


def mask_value(name, value):
    value = "" if value is None else str(value)
    if SECRET.search(name or ""):
        return "***MASKED***"
    if len(value) > 60 and re.fullmatch(r"[A-Za-z0-9\-_.=+/]+", value):
        return "***MASKED (long token-like value)***"
    return value


def mask_url(url):
    if "?" not in url:
        return url
    base, query = url.split("?", 1)
    parts = []
    for p in query.split("&"):
        if "=" in p:
            k, v = p.split("=", 1)
            parts.append(f"{k}={mask_value(k, v)}")
        else:
            parts.append(p)
    return base + "?" + "&".join(parts)


def mask_body(body):
    if not body:
        return ""
    body = re.sub(r'("([^"]+)"\s*:\s*)"([^"]*)"',
                  lambda m: m.group(1) + '"***MASKED***"' if SECRET.search(m.group(2)) else m.group(0), body)
    body = re.sub(r"([A-Za-z0-9_\-.]+)=([^&\s]*)",
                  lambda m: m.group(1) + "=***MASKED***" if SECRET.search(m.group(1)) else m.group(0), body)
    return body[:3000] + (" ...(truncated)" if len(body) > 3000 else "")


def header(headers, name):
    for h in headers or []:
        if str(h.get("name", "")).lower() == name.lower():
            return str(h.get("value", ""))
    return ""


def search_dirs():
    home = os.path.expanduser("~")
    dirs = [os.path.join(home, "Desktop"), os.path.join(home, "Downloads"),
            os.path.dirname(os.path.abspath(__file__))]
    dirs += glob.glob(os.path.join(home, "OneDrive*", "Desktop"))
    return [d for d in dict.fromkeys(dirs) if os.path.isdir(d)]


def newest(pattern_ok):
    best = None
    for d in search_dirs():
        for name in os.listdir(d):
            path = os.path.join(d, name)
            if os.path.isfile(path) and not name.startswith("~$") and pattern_ok(name.lower()):
                if best is None or os.path.getmtime(path) > os.path.getmtime(best):
                    best = path
    return best


def desktop():
    for d in [os.path.join(os.path.expanduser("~"), "Desktop")] + glob.glob(
            os.path.join(os.path.expanduser("~"), "OneDrive*", "Desktop")):
        if os.path.isdir(d):
            return d
    return os.path.dirname(os.path.abspath(__file__))


def shape(v):
    v = "" if v is None else str(v).strip()
    if not v:
        return "(empty)"
    if re.fullmatch(r"\d+(\.0)?", v):
        return f"digits x{len(v.replace('.0', ''))}"
    if re.match(r"\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}", v):
        return "date"
    if re.fullmatch(r"[A-Z0-9\-]+", v):
        return f"code x{len(v)}"
    return f"text x{len(v)}"


def main():
    skip_download = "--skip-download-test" in sys.argv
    out("=== TracIT EUC report - export details (secrets masked) ===")
    out(f"Collected: {datetime.now():%Y-%m-%d %H:%M}   Python {sys.version.split()[0]}")
    out(f"Looked in: {', '.join(search_dirs())}")
    out("")

    # ---------------------------------------------------- 1. HAR file
    candidates = []
    har_path = newest(lambda n: n.endswith(".har"))
    if not har_path:
        out("[1] EXPORT REQUEST: no .har file found on the Desktop / in Downloads.")
        out("    Edge: F12 -> Network -> click Export in TracIT -> 'Export HAR' button -> save, then re-run.")
    else:
        out(f"[1] EXPORT REQUEST  (from {os.path.basename(har_path)})")
        with open(har_path, encoding="utf-8-sig") as f:
            entries = json.load(f).get("log", {}).get("entries", [])
        out(f"    {len(entries)} requests in the log")
        for i, e in enumerate(entries):
            url = e.get("request", {}).get("url", "")
            if STATIC.search(url) or url.startswith("data:"):
                continue
            ctype = header(e.get("response", {}).get("headers"), "content-type")
            disp = header(e.get("response", {}).get("headers"), "content-disposition")
            score = 0
            if re.search(r"spreadsheet|excel|csv|octet-stream", ctype, re.I):
                score += 3
            if re.search(r"attachment|filename", disp, re.I):
                score += 3
            if re.search(r"export|download|report|excel|csv|xlsx", url, re.I):
                score += 2
            if score:
                candidates.append((score, i, e, ctype, disp))
        candidates.sort(key=lambda c: -c[0])
        candidates = candidates[:4]
        if not candidates:
            out("    Could not spot an export/download request (was 'Preserve log' on, and did the file finish?).")
            out("    Last 10 non-static requests, for reference:")
            for e in [x for x in entries if not STATIC.search(x.get("request", {}).get("url", ""))][-10:]:
                out(f"      {e['request'].get('method')} {mask_url(e['request'].get('url', ''))} -> "
                    f"{e.get('response', {}).get('status')}")
        for n, (score, i, e, ctype, disp) in enumerate(candidates, 1):
            req, resp = e.get("request", {}), e.get("response", {})
            out("")
            out(f"  --- Candidate {n} (match score {score}) ---")
            out(f"    Method      : {req.get('method')}")
            out(f"    URL         : {mask_url(req.get('url', ''))}")
            out(f"    Status      : {resp.get('status')}")
            out(f"    Content-Type: {ctype}")
            out(f"    Disposition : {disp}")
            out(f"    Size (bytes): {resp.get('content', {}).get('size')}")
            out("    Request headers (secret values masked):")
            for h in req.get("headers", []):
                if not str(h.get("name", "")).startswith(":"):
                    out(f"      {h.get('name')}: {mask_value(h.get('name'), h.get('value'))}")
            if req.get("postData"):
                out(f"    Payload ({req['postData'].get('mimeType')}):")
                out(f"      {mask_body(req['postData'].get('text', ''))}")
        if candidates:
            idx = candidates[0][1]
            out("")
            out("    Requests just before the export (to spot a two-step export):")
            for p in entries[max(0, idx - 6):idx]:
                if STATIC.search(p.get("request", {}).get("url", "")):
                    continue
                out(f"      {p['request'].get('method')} {mask_url(p['request'].get('url', ''))} -> "
                    f"{p.get('response', {}).get('status')} {header(p.get('response', {}).get('headers'), 'content-type')}")
    out("")

    # -------------------------------------- 2. direct download test
    out("[2] DOES THE LINK WORK WITH ONLY YOUR WINDOWS LOGIN?  (no browser cookies - same as the app)")
    if skip_download:
        out("    Skipped (--skip-download-test).")
    elif not candidates:
        out("    Skipped - no export request found in step 1.")
    else:
        e = candidates[0][2]
        req = e.get("request", {})
        try:
            import requests
            try:
                from requests_negotiate_sspi import HttpNegotiateAuth
                auth, auth_name = HttpNegotiateAuth(), "Windows login (requests-negotiate-sspi)"
            except ImportError:
                auth, auth_name = None, "NO Windows login (requests-negotiate-sspi not installed)"
            try:
                import truststore
                truststore.inject_into_ssl()
            except Exception:
                pass
            kwargs = {"auth": auth, "timeout": 120, "allow_redirects": True}
            if req.get("postData", {}).get("text"):
                kwargs["data"] = req["postData"]["text"].encode("utf-8")
                kwargs["headers"] = {"Content-Type": req["postData"].get("mimeType") or "application/json"}
            r = requests.request(req.get("method", "GET"), req.get("url"), **kwargs)
            body = r.content
            if body[:2] == b"PK":
                kind = "XLSX (Excel file) - GOOD"
            else:
                head = body[:400].decode("utf-8", "replace")
                if re.search(r"<html|<!doctype", head, re.I):
                    kind = "HTML page (probably a LOGIN page) - needs browser sign-in"
                elif head.lstrip()[:1] in ("{", "["):
                    kind = "JSON (maybe a job id / link - two-step export)"
                elif "," in head:
                    kind = "CSV text - GOOD"
                else:
                    kind = "unknown"
            out(f"    Auth used : {auth_name}")
            out(f"    HTTP {r.status_code}, Content-Type: {r.headers.get('Content-Type')}, {len(body)} bytes"
                + (f", redirected via {len(r.history)} hop(s)" if r.history else ""))
            out(f"    Result    : {kind}")
            if kind.startswith("JSON"):
                out(f"    JSON start (masked): {mask_body(body[:500].decode('utf-8', 'replace'))}")
        except Exception as exc:
            out(f"    FAILED: {type(exc).__name__}: {exc}")
            out("    (401/403/login redirect = TracIT needs the browser session -> the app will use its hidden window.)")
    out("")

    # ------------------------------------------------- 3. report file
    out("[3] REPORT FILE (headers + row count only - no data values)")
    report = newest(lambda n: "euc" in n and n.endswith((".xlsx", ".xlsm", ".csv")))
    if not report:
        out("    No EUC report found on the Desktop / in Downloads (file name containing 'euc').")
    else:
        out(f"    File: {os.path.basename(report)}  ({os.path.getsize(report) // 1024} KB, "
            f"saved {datetime.fromtimestamp(os.path.getmtime(report)):%Y-%m-%d %H:%M})")
        rows, total = [], 0
        try:
            if report.endswith(".csv"):
                import csv
                for enc in ("utf-8-sig", "cp1252"):
                    try:
                        with open(report, newline="", encoding=enc) as f:
                            allrows = list(csv.reader(f))
                        break
                    except UnicodeDecodeError:
                        continue
                rows, total = allrows[:12], len(allrows)
            else:
                import openpyxl
                wb = openpyxl.load_workbook(report, read_only=True, data_only=True)
                ws = wb[wb.sheetnames[0]]
                for r in ws.iter_rows(values_only=True):
                    total += 1
                    if total <= 12:
                        rows.append(["" if v is None else str(v) for v in r])
                wb.close()
        except Exception as exc:
            out(f"    Could not read it ({exc}). Send a screenshot of the header row instead.")
        if rows:
            hi = next((i for i, r in enumerate(rows) if sum(1 for v in r if str(v).strip()) >= 3), 0)
            sample = rows[hi + 1] if len(rows) > hi + 1 else []
            out(f"    Header row is row {hi + 1}; about {total - hi - 1} data rows")
            out("    Columns (with the SHAPE of the first data row - values are not copied):")
            for c, h in enumerate(rows[hi]):
                if str(h).strip():
                    out(f"      {c + 1:2d}. {str(h):35s} {shape(sample[c]) if c < len(sample) else ''}")
    out("")
    out("=== end ===")

    dest = os.path.join(desktop(), "TracIT_Export_Details.txt")
    with open(dest, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"\nSaved to: {dest}\nOpen it, check you're happy with what it contains, then paste it to Claude.")


if __name__ == "__main__":
    main()
